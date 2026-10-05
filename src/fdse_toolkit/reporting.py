"""Enterprise FDSE report bundle generator for validated engagement data."""
 
from __future__ import annotations
 
import hashlib
import html
import json
import re
from collections.abc import Iterable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
 
from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
 
from .contracts import validate_document
from .risk import correlate_findings
 
 
def _slug(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._-")
    return value[:80] or "engagement"
 
 
def _safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value
 
 
def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()
 
 
def _validate_engagement(data: dict[str, Any]) -> dict[str, Any]:
    validate_document("engagement", data)
    for finding_id in data.get("finding_ids", []):
        # Full finding objects are accepted by generate_report_bundle; IDs are
        # only relationship references and are not resolved here.
        if not isinstance(finding_id, str):
            raise TypeError("finding_ids must contain strings")
    return data
 
 
def _severity_counts(findings: Iterable[dict[str, Any]]) -> dict[str, int]:
    counts = {key: 0 for key in ("CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO")}
    for finding in findings:
        severity = str(finding["severity"]).upper()
        counts[severity] = counts.get(severity, 0) + 1
    return counts
 
 
def build_pdf(engagement: dict[str, Any], findings: list[dict[str, Any]], output: Path) -> None:
    client = engagement["client"]["name"]
    counts = _severity_counts(findings)
    doc = SimpleDocTemplate(
        str(output), pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
        topMargin=1.8 * cm, bottomMargin=1.8 * cm,
        title=f"Tinlance FDSE Security Assessment — {client}",
        author="Tinlance Limited",
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("TINLANCE FDSE SECURITY ASSESSMENT", styles["Title"]),
        Paragraph(html.html.escape(client), styles["Heading2"]),
        Paragraph(escape(f"Engagement: {engagement['engagement_id']}"), styles["Normal"]),
        Paragraph(escape(f"Classification: {engagement.get('data_classification', 'CONFIDENTIAL')}"), styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Executive Summary", styles["Heading1"]),
        Paragraph(f"Validated findings: {len(findings)}", styles["Normal"]),
    ]
    summary = [["Severity", "Count"]] + [[k, str(v)] for k, v in counts.items()]
    table = Table(summary, colWidths=[6 * cm, 3 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b6b61")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ALIGN", (1, 1), (1, -1), "CENTER"),
    ]))
    story.extend([table, Spacer(1, 14), Paragraph("Findings", styles["Heading1"])])
    for finding in findings:
        story.extend([
            Paragraph(escape(f"{finding['severity']} — {finding['title']}"), styles["Heading2"]),
            Paragraph(escape(f"Finding ID: {finding['finding_id']}"), styles["Normal"]),
            Paragraph(escape(f"Confidence: {finding['confidence']:.0%}"), styles["Normal"]),
            Paragraph(escape(f"Evidence: {', '.join(finding['evidence_ids'])}"), styles["Normal"]),
            Paragraph(escape(finding.get("description", "") or "No description provided."), styles["BodyText"]),
            Spacer(1, 8),
        ])
    doc.build(story)
 
 
def build_docx(engagement: dict[str, Any], findings: list[dict[str, Any]], output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    doc.add_heading("Tinlance FDSE Security Assessment", 0)
    doc.add_paragraph(engagement["client"]["name"])
    doc.add_paragraph(f"Engagement ID: {engagement['engagement_id']}")
    doc.add_heading("Executive Summary", 1)
    doc.add_paragraph(f"Validated findings: {len(findings)}")
    table = doc.add_table(rows=1, cols=5)
    for cell, text in zip(table.rows[0].cells, ["ID", "Severity", "Title", "Confidence", "Evidence"]):
        cell.text = text
    for finding in findings:
        cells = table.add_row().cells
        cells[0].text = finding["finding_id"]
        cells[1].text = finding["severity"]
        cells[2].text = finding["title"]
        cells[3].text = f"{finding['confidence']:.0%}"
        cells[4].text = ", ".join(finding["evidence_ids"])
    doc.add_heading("Scope", 1)
    for target in engagement["scope"]["authorized_targets"]:
        doc.add_paragraph(target, style="List Bullet")
    doc.add_heading("Technical Findings", 1)
    for finding in findings:
        doc.add_heading(f"{finding['severity']} — {finding['title']}", 2)
        doc.add_paragraph(finding.get("description", "") or "No description provided.")
        doc.add_paragraph(f"Evidence: {', '.join(finding['evidence_ids'])}")
        if finding.get("remediation_ids"):
            doc.add_paragraph(f"Remediation references: {', '.join(finding['remediation_ids'])}")
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(10)
    doc.save(output)
 
 
def build_excel(engagement: dict[str, Any], findings: list[dict[str, Any]], output: Path) -> None:
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"
    dashboard["A1"] = "TINLANCE FDSE SECURITY ASSESSMENT"
    dashboard["A2"] = _safe_cell(engagement["client"]["name"])
    dashboard["A3"] = _safe_cell(engagement["engagement_id"])
    dashboard["A5"] = "Severity"
    dashboard["B5"] = "Count"
    for row, (severity, count) in enumerate(_severity_counts(findings).items(), start=6):
        dashboard.cell(row, 1, severity)
        dashboard.cell(row, 2, count)
    sheet = wb.create_sheet("Findings")
    headers = ["Finding ID", "Severity", "Title", "Confidence", "Risk Score", "Status", "Evidence IDs", "Sources"]
    for col, header in enumerate(headers, start=1):
        sheet.cell(1, col, header)
    for row, finding in enumerate(findings, start=2):
        values = [
            finding["finding_id"], finding["severity"], finding["title"], finding["confidence"],
            finding.get("risk_score"), finding["status"], ", ".join(finding["evidence_ids"]),
            ", ".join(finding.get("correlated_sources", [])),
        ]
        for col, value in enumerate(values, start=1):
            sheet.cell(row, col, _safe_cell(value))
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0B6B61")
        ws.freeze_panes = "A2"
        for column in ws.columns:
            letter = column[0].column_letter
            ws.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in column) + 2, 12), 50)
    wb.save(output)
 
 
def generate_report_bundle(
    engagement: dict[str, Any],
    findings: list[dict[str, Any]],
    output_dir: Path,
    *,
    formats: tuple[str, ...] = ("pdf", "docx", "xlsx"),
) -> dict[str, Any]:
    """Validate inputs and generate board, technical, and operational reports."""
    _validate_engagement(engagement)
    normalized = correlate_findings(findings)
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    client_slug = _slug(engagement["client"]["name"])
    generated_at = datetime.now(UTC).isoformat().replace("+00:00", "Z")
    artifacts: list[dict[str, Any]] = []
    builders = {"pdf": build_pdf, "docx": build_docx, "xlsx": build_excel}
    suffixes = {"pdf": ".pdf", "docx": ".docx", "xlsx": ".xlsx"}
    for fmt in formats:
        if fmt not in builders:
            raise ValueError(f"unsupported report format: {fmt}")
        path = out / f"TinlanceReport_{client_slug}_{engagement['engagement_id']}{suffixes[fmt]}"
        builders[fmt](engagement, normalized, path)
        artifacts.append({"name": path.name, "sha256": _sha256(path), "size_bytes": path.stat().st_size, "format": fmt})
    manifest = {
        "schema_version": "1.0.0",
        "engagement_id": engagement["engagement_id"],
        "generated_at": generated_at,
        "artifacts": artifacts,
        "finding_count": len(normalized),
    }
    validate_document("report-manifest", manifest)
    (out / "report_manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return manifest

"""Client-specific incident-response playbook generator."""
from __future__ import annotations

from pathlib import Path

from docx import Document

from .contracts import validate_document

PLAYBOOKS = {
    "C2": ("C2 / Beaconing", ["Preserve volatile evidence", "Isolate affected assets under the approved response plan", "Block confirmed malicious infrastructure", "Scope lateral movement and exfiltration", "Document containment and recovery decisions"]),
    "CREDENTIAL_THEFT": ("Credential Theft", ["Preserve the exposure evidence", "Revoke or disable affected credentials", "Rotate dependent secrets and sessions", "Audit credential use and privilege changes", "Verify removal from source and build history"]),
    "LATERAL_MOVEMENT": ("Lateral Movement", ["Preserve authentication and endpoint telemetry", "Contain affected network segments", "Identify compromised accounts and paths", "Reset credentials according to incident policy", "Verify segmentation and persistence removal"]),
    "RANSOMWARE": ("Ransomware", ["Preserve evidence and volatile state", "Contain affected systems and protect clean backups", "Identify patient zero and propagation path", "Coordinate executive, legal, and recovery decisions", "Restore only from verified clean recovery points"]),
    "SUPPLY_CHAIN": ("Supply-Chain Compromise", ["Freeze affected dependency or vendor change", "Preserve package/build evidence", "Identify impacted releases and environments", "Rotate credentials and signing material if exposure is possible", "Verify clean rebuild and deployment"]),
    "AI_AGENT_ABUSE": ("AI Agent Abuse", ["Preserve prompts, tool calls, approvals, and evidence references", "Disable or constrain affected agent capabilities", "Review authorization and policy decisions", "Identify data/tool impact", "Revalidate controls before restoring capability"]),
}


def generate_playbook(engagement: dict, incident: dict, output: Path) -> Path:
    validate_document("engagement", engagement)
    validate_document("incident", incident)
    if incident["category"] not in PLAYBOOKS:
        raise ValueError("unsupported incident category")
    title, actions = PLAYBOOKS[incident["category"]]
    doc = Document()
    doc.add_heading("Tinlance FDSE Incident Response Playbook", 0)
    doc.add_paragraph(engagement["client"]["name"])
    doc.add_paragraph(f"Incident: {incident['incident_id']} | Category: {title} | Severity: {incident['severity']}")
    doc.add_heading("Trigger and Classification", 1)
    doc.add_paragraph("Use this playbook only after the incident is classified and the engagement authorization/incident-response authority is confirmed.")
    doc.add_heading("Immediate Response", 1)
    for action in actions:
        doc.add_paragraph(action, style="List Number")
    doc.add_heading("Evidence Preservation", 1)
    doc.add_paragraph("Record every material action, actor, timestamp, source artifact, and evidence identifier. Preserve original evidence before analysis where practical.")
    doc.add_heading("Communications", 1)
    doc.add_paragraph("Notify the customer-defined incident lead, executive stakeholders, legal/privacy contacts, and other authorized parties according to the customer's incident-response plan.")
    doc.add_heading("Regulatory Assessment", 1)
    doc.add_paragraph("Determine which regulations apply to this entity, sector, jurisdiction, and incident classification. Record the authoritative source, effective date, competent authority, and applicable deadline. Do not infer a legal deadline from this template.")
    doc.add_heading("Recovery and Closure", 1)
    doc.add_paragraph("Document eradication, recovery validation, residual risk, lessons learned, retest requirements, and evidence-retention/deletion decisions.")
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output

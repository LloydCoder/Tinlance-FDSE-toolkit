from pathlib import Path

from openpyxl import load_workbook
from docx import Document

from fdse_toolkit.reporting import generate_report_bundle


def test_report_bundle_generates_three_valid_artifacts(tmp_path: Path):
    engagement = {
        "schema_version": "1.0.0",
        "engagement_id": "FDSE-ACME-001",
        "client": {"name": "Acme <Security>"},
        "scope": {"authorized_targets": ["example.com"]},
        "created_at": "2026-10-05T08:00:00Z",
        "data_classification": "CONFIDENTIAL",
    }
    findings = [{
        "schema_version": "1.0.0", "finding_id": "FND-AAA", "title": "Injected <title>",
        "description": "Synthetic <b>description</b>", "severity": "HIGH", "status": "OPEN",
        "confidence": 0.9, "evidence_ids": ["EVD-AAA"], "asset_ids": [], "source": "fixture",
    }]
    manifest = generate_report_bundle(engagement, findings, tmp_path)
    assert manifest["finding_count"] == 1
    assert len(manifest["artifacts"]) == 3
    pdf = next(tmp_path.glob("*.pdf"))
    docx = next(tmp_path.glob("*.docx"))
    xlsx = next(tmp_path.glob("*.xlsx"))
    assert pdf.read_bytes().startswith(b"%PDF-")
    assert Document(docx).paragraphs
    assert "Findings" in load_workbook(xlsx, read_only=True).sheetnames
    assert (tmp_path / "report_manifest.json").is_file()
